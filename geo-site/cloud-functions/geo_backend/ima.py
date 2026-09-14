from __future__ import annotations

import hmac
import json
import re
import asyncio
import uuid
from io import BytesIO
from datetime import date, datetime, timezone
from urllib.parse import urlparse

import httpx

from .errors import ApiError
from .security import digest


def normalize(value: object) -> str:
    return re.sub(r"\s+", "", str(value or "")).casefold()


class ImaCache:
    """Site-wide IMA cache facade backed by the repository and cache generation."""

    def __init__(self, repository: object, master_key: str = "", *, generation: int | None = None) -> None:
        self.repository = repository
        self.master_key = master_key
        self.generation = generation

    @staticmethod
    def key(kind: str, request: dict[str, object], generation: int = 1) -> str:
        normalized = {
            str(name): re.sub(r"\s+", " ", value.strip()) if name == 'query' and isinstance(value, str) else value
            for name, value in sorted(request.items())
        }
        return digest(str(generation) + ":" + kind + ":" + json.dumps(normalized, ensure_ascii=False, sort_keys=True, separators=(",", ":")))

    async def get_or_fetch(
        self,
        kind: str,
        request: dict[str, object],
        fetch,
    ) -> object:
        generation = self.generation if self.generation is not None else int(self.repository.get_ima_cache_generation())
        cache_key = self.key(kind, request, generation)
        cached = self.repository.get_ima_cache(kind, cache_key, generation, self.master_key)
        if cached is not None:
            return cached
        acquire = getattr(self.repository, "acquire_ima_cache_lock", None)
        release = getattr(self.repository, "release_ima_cache_lock", None)
        owner_token = uuid.uuid4().hex
        owns_lock = False
        if acquire and release:
            for _ in range(121):
                owns_lock = bool(acquire(cache_key, generation, owner_token, 150))
                if owns_lock:
                    break
                await asyncio.sleep(0.25)
                cached = self.repository.get_ima_cache(kind, cache_key, generation, self.master_key)
                if cached is not None:
                    return cached
            if not owns_lock:
                raise ApiError(503, "IMA 缓存正在由其他任务更新，请稍后继续。", "IMA_CACHE_BUSY")
        try:
            # Another request may have filled the cache between our miss and lock acquisition.
            cached = self.repository.get_ima_cache(kind, cache_key, generation, self.master_key)
            if cached is not None:
                return cached
            value = await fetch()
            if value is None:
                raise ApiError(502, "IMA 返回空内容，未写入缓存。", "IMA_EMPTY")
            self.repository.put_ima_cache(kind, cache_key, generation, value, request, self.master_key)
            return value
        finally:
            if owns_lock:
                release(cache_key, owner_token)

    def clear_generation(self, user_id: str | None = None) -> int:
        return int(self.repository.clear_ima_cache_generation(user_id))


def load_ima_credentials(repository: object, master_key: str, client_id: str, api_key: str) -> dict[str, object]:
    stored = repository.load_ima(master_key)
    return stored or {"clientId": client_id, "apiKey": api_key, "expiresAt": "", "updatedAt": None}


async def ima_post(credentials: dict[str, object], path: str, payload: dict[str, object], *, client=None) -> dict[str, object]:
    if not credentials.get("clientId") or not credentials.get("apiKey"):
        raise ApiError(503, "共享 IMA 凭据尚未配置，请联系管理员。", "IMA_REQUIRED")
    owned = client is None
    http = client or httpx.AsyncClient(timeout=httpx.Timeout(25.0), follow_redirects=False)
    try:
        response = await http.post(
            f"https://ima.qq.com/{path}",
            headers={
                "Content-Type": "application/json; charset=utf-8",
                "ima-openapi-clientid": str(credentials["clientId"]),
                "ima-openapi-apikey": str(credentials["apiKey"]),
            },
            json=payload,
        )
    except httpx.TransportError:
        raise ApiError(502, "IMA 连接失败，未自动重试。", "IMA_NETWORK")
    finally:
        if owned:
            await http.aclose()
    if not response.is_success:
        raise ApiError(502, f"IMA 请求失败（HTTP {response.status_code}）。", "IMA_ERROR")
    try:
        result = response.json()
    except ValueError:
        raise ApiError(502, "IMA 返回格式错误。", "IMA_ERROR")
    if result.get("code") != 0 or not isinstance(result.get("data"), dict):
        raise ApiError(502, f"IMA 凭据、权限或请求异常（代码 {result.get('code', '未知')}），请检查有效期。", "IMA_ERROR")
    return result["data"]


async def update_ima_credentials(
    repository: object,
    master_key: str,
    admin_secret: str,
    body: dict[str, object],
    *,
    client=None,
) -> dict[str, object]:
    now = datetime.now(timezone.utc)
    if not repository.take_admin_attempt(digest("ima-admin-update"), now, 5):
        raise ApiError(429, "管理员更新尝试过多，请稍后再试。")
    if len(admin_secret) < 24:
        raise ApiError(503, "管理员更新口令尚未配置。", "ADMIN_SETUP_REQUIRED")
    supplied = str(body.get("adminSecret", ""))
    if not hmac.compare_digest(supplied, admin_secret):
        raise ApiError(403, "管理员更新口令不正确。", "ADMIN_REQUIRED")
    value = {
        "clientId": str(body.get("clientId", "")).strip(),
        "apiKey": str(body.get("apiKey", "")).strip(),
        "expiresAt": str(body.get("expiresAt", "")).strip(),
        "updatedAt": now.isoformat(),
    }
    if not value["clientId"] or not value["apiKey"] or len(value["clientId"]) > 4096 or len(value["apiKey"]) > 4096:
        raise ApiError(400, "IMA 凭据未填写或过长。")
    if any(char in value["clientId"] + value["apiKey"] for char in "\r\n\x00"):
        raise ApiError(400, "IMA 凭据格式不正确。")
    try:
        expiry = date.fromisoformat(value["expiresAt"])
    except ValueError:
        raise ApiError(400, "请选择有效的未来到期日期。")
    if expiry < now.date():
        raise ApiError(400, "请选择有效的未来到期日期。")
    data = await ima_post(value, "openapi/wiki/v1/search_knowledge_base", {"query": "copilot", "cursor": "", "limit": 20}, client=client)
    if not any(normalize(item.get("name") or item.get("kb_name")) == "copilot" for item in data.get("info_list", [])):
        raise ApiError(422, "新凭据无法访问 copilot 知识库，已保留旧凭据。")
    repository.save_ima(value, master_key)
    return {"saved": True, "expiresAt": value["expiresAt"]}


def next_cursor(data: dict[str, object], current: str = "") -> str | None:
    if data.get("is_end") is True:
        return None
    value = data.get("next_cursor")
    if not isinstance(value, str) or not value or value == current:
        raise ApiError(502, "IMA 分页信息异常，已停止以避免重复读取。")
    return value


async def read_media(credentials: dict[str, object], media: dict[str, object], *, client=None) -> dict[str, str]:
    info = await ima_post(credentials, "openapi/wiki/v1/get_media_info", {"media_id": media["media_id"]}, client=client)
    if info.get("media_type") == 11 and (info.get("notebook_ext_info") or {}).get("notebook_id"):
        note = await ima_post(
            credentials,
            "openapi/note/v1/get_doc_content",
            {"note_id": info["notebook_ext_info"]["notebook_id"], "target_content_format": 0},
            client=client,
        )
        text = note.get("content")
    else:
        url = str((info.get("url_info") or {}).get("url", ""))
        parsed = urlparse(url)
        hostname = (parsed.hostname or "").lower()
        try:
            port = parsed.port
        except ValueError:
            port = -1
        if (
            parsed.scheme != "https"
            or parsed.username
            or parsed.password
            or (port not in {None, 443})
            or not (hostname == "ima.qq.com" or hostname.endswith(".myqcloud.com"))
        ):
            raise ApiError(422, f"IMA 原文下载域名不支持：{media['title']}。请在 IMA 上传文件版本。")
        headers = (info.get("url_info") or {}).get("headers") or {}
        if not isinstance(headers, dict) or any(re.search(r"ima-openapi|cookie|host", str(key), re.IGNORECASE) for key in headers):
            raise ApiError(502, "IMA 下载响应包含不允许的请求头。")
        owned = client is None
        http = client or httpx.AsyncClient(timeout=httpx.Timeout(25.0), follow_redirects=False)
        try:
            response = await http.get(url, headers={str(key): str(value) for key, value in headers.items()})
        except httpx.TransportError:
            raise ApiError(502, f"IMA 原文下载失败：{media['title']}。")
        finally:
            if owned:
                await http.aclose()
        if not response.is_success:
            raise ApiError(502, f"IMA 原文下载失败：{media['title']}。")
        content_length = int(response.headers.get("content-length", "0") or 0)
        if content_length > 20 * 1024 * 1024 or len(response.content) > 20 * 1024 * 1024:
            raise ApiError(413, f"IMA 原文超过 20 MB：{media['title']}。")
        content = response.content
        title = str(media["title"])
        if content[:4] == b"%PDF":
            from pypdf import PdfReader

            pdf = PdfReader(BytesIO(content))
            if len(pdf.pages) > 100:
                raise ApiError(413, f"IMA PDF 超过 100 页，请拆分：{title}。")
            text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        elif content[:2] == b"PK" and title.lower().endswith(".docx"):
            from docx import Document

            document = Document(BytesIO(content))
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                parts.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
            text = "\n".join(parts)
        elif re.search(r"\.(txt|md|markdown)$", title, re.IGNORECASE):
            try:
                text = content.decode("utf-8", errors="strict")
            except UnicodeDecodeError:
                text = content.decode("gb18030", errors="strict")
        else:
            raise ApiError(422, f"IMA 原文格式暂不支持：{title}。请转为 DOCX、PDF、TXT 或 MD。")
    if not isinstance(text, str) or not text.strip():
        raise ApiError(422, f"IMA 原文为空或为扫描图片：{media['title']}。")
    if len(text) > 180000:
        raise ApiError(413, f"IMA 原文过长，请拆分：{media['title']}。未使用截断内容。")
    return {"title": str(media["title"]), "text": text.strip()}
